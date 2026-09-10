import json
import pandas as pd
from openpyxl import load_workbook
from datetime import datetime
import re

EXCEL_PATH = r'c:\dev\ecopin_image_validation\Ecopin Gantt Chart.xlsx'
JSON_PATH = r'c:\dev\ecopin_image_validation\scraped_metadata.json'
SHEET_NAME = 'Dataset'

DOWNLOAD_DATE = datetime(2026, 9, 9)

PRIMARY_LABEL_MAP = {
    'WST': 'waste',
    'FLD': 'flooding',
    'POL': 'pollution',
    'NEG': 'non_environmental',
}

WASTE_SUBCATEGORY_RULES = [
    ('roadside_dumping', [
        r'roadside', r'road.?side', r'fly.?tip', r'fly.?tipping', r'illeg.?dump',
        r'dumped.*road', r'road.*dump', r'dumping', r'curb.?side', r'kerb.?side',
        r'street.*dump', r'dump.*street', r'layby', r'lay.?by',
    ]),
    ('waterway_waste', [
        r'waterway', r'river.*garbage', r'river.*trash', r'river.*waste', r'river.*litter',
        r'ocean.*garbage', r'ocean.*trash', r'sea.*trash', r'sea.*garbage',
        r'beach.*litter', r'beach.*trash', r'beach.*waste', r'shore.*trash', r'shore.*waste',
        r'coast.*trash', r'coast.*waste', r'harbour.*trash', r'harbor.*trash',
        r'plastic.*ocean', r'plastic.*sea', r'ship.*garbage', r'water.*trash',
        r'water.*waste', r'canal.*trash', r'canal.*waste', r'water.?pollut.*trash',
        r'reef.*trash', r'wetland.*trash', r'marine.*debris', r'marine.*litter',
    ]),
    ('garbage_accumulation', [
        r'garbage.?accumul', r'pile.*garbage', r'pile.*trash', r'pile.*rubbish',
        r'mountain.*garbage', r'mountain.*trash', r'mound.*trash', r'heaps?.*trash',
        r'heaps?.*garbage', r'heaps?.*rubbish', r'overflow.*bin', r'overflow.*trash',
        r'overflow.*garbage', r'landfill', r'garbage.?dump', r'trash.?dump',
        r'rubbish.?dump', r'waste.?dump', r'dump.?site', r'compost',
        r'bin.?strike', r'garbage.?collect', r'waste.?collect', r'trash.?collect',
        r'collect.*garbage', r'collect.*trash', r'garbage.?bags?', r'trash.?bags?',
        r'rubbish.?bags?', r'burn.?rubbish', r'burn.?trash', r'burn.?garbage',
        r'incinerat', r'smouldering.?trash', r'smoldering.?trash', r'garbage.?burn',
        r'vacant.?lot.*litt', r'vacant.?lot.*trash', r'vacant.?lot.*garbage',
        r'garbage', r'trash', r'rubbish', r'litter', r'detritus', r'debris',
        r'waste', r'recycl', r'clean.?up', r'cleanup',
    ]),
    ('litter_accumulation', [
        r'litter', r'scatter.*trash', r'scatter.?litt', r'fast.?food',
        r'eat.?and.?run', r'food.*packag', r'discarded', r'street.*litt',
        r'sidewalk.*litt', r'pavement.*litt', r'campsite.*litt', r'ground.*litt',
        r'park.?litt', r'school.?litt', r'plaza.?litt', r'public.*litt',
    ]),
    ('public_area_waste', [
        r'park.*garbage', r'park.*waste', r'park.*trash', r'plaza.*trash',
        r'market.*waste', r'market.?garbage', r'campsite', r'festival.*trash',
        r'stadium.*waste', r'beac', r'school.*waste', r'hospital.*waste',
        r'station.*trash', r'station.*waste', r'terminal.*waste',
    ]),
    ('electronics', [
        r'electronic', r'e.?waste', r'ewaste', r'e.?scrap', r'circuit.?board',
        r'computer.*waste', r'appliance.*waste', r'electronic.?scrap',
        r'pcb.*waste', r'battery.*waste', r'mobile.*waste', r'phone.*waste',
        r'cathode.*ray', r'crt.*waste', r'electronic.?dum',
    ]),
]

FLOODING_SUBCATEGORY_RULES = [
    ('flooded_road', [
        r'flood.*road', r'road.*flood', r'street.*flood', r'flood.*street',
        r'traffic.*flood', r'car.*flood', r'vehicle.*flood', r'highway.*flood',
        r'road.*water', r'street.*water', r'flooded.*street', r'flooded.*road',
        r'flooded.*highway', r'under.?water.*road', r'road.?submerg',
        r'street.?submerg', r'floodway', r'flash.?flood.*road', r'flood.*car',
        r'banjir.*jalan', r'inunda.*rua', r'inunda.*via', r'inunda.?calle',
        r'swimming.?pool', r'canal.?overflow.*road', r'flood.*thai',
    ]),
    ('flooded_residential_area', [
        r'flood.*house', r'house.*flood', r'flood.*home', r'home.*flood',
        r'flood.*resident', r'resident.*flood', r'flood.*villa', r'flood.*neighbor',
        r'neighbor.*flood', r'flood.?suburb', r'suburb.*flood',
        r'flood.*market', r'flooded.?market', r'flood.*village', r'village.*flood',
        r'flood.*kampung', r'flood.?town', r'town.*flood', r'flood.*city',
        r'city.*flood', r'flood.?urba', r'urban.?flood', r'house.?affect',
        r'home.?affect', r'carport', r'rumah', r'perumahan', r'permukiman',
        r'residential.?flo', r'flooded.?house', r'flooded.?home',
        r'flood.?district', r'flood.?locality', r'flooded.?area.*house',
        r'flood.?estate', r'flooded.?neighbor', r'rice.*flood', r'rice.?padd',
        r'swamp', r'flooded.?courtyard', r'courtyard.?flood',
    ]),
    ('urban_flooding', [
        r'urban.?flood', r'flood.?urba', r'city.?flood', r'capital.?flood',
        r'metro.?flood', r'municipal.?flood', r'business.?district.*flood',
        r'flood.*town.*center', r'town.*center.*flood', r'flood.*jakarta',
        r'flood.*bangkok', r'flood.*hanoi', r'flood.*manila', r'flood.*mumbai',
        r'flood.*chennai', r'jakarta.*banjir', r'banjir.?jakarta',
        r'flood.*vietnam', r'flood.*thailand', r'bangkok.?flood',
        r'urban.?swamp', r'water.?log', r'waterlog', r'logging.*road',
        r'logging.?street',
    ]),
    ('river_overflow', [
        r'river.?overflow', r'overflow.*river', r'river.*burst', r'burst.*river',
        r'river.?flood', r'flood.*river', r'riparian.?flood', r'fluvial.?flood',
        r'levee.*breach', r'embankment.*breach', r'dam.?overflow', r'weir.*overflow',
        r'river.*rise', r'rising.?water.*river', r'river.?submerg',
        r'sungai', r'chao.?phraya', r'mekong', r'ganges.*flood', r'irrawaddy',
        r'emajogi', r'citarum', r'river.?level', r'water.?course.?overflow',
    ]),
    ('drainage_flooding', [
        r'drain.*block', r'block.*drain', r'drainage.?poor', r'poor.?drain',
        r'sewer.*overflow', r'sewage.?overflow', r'storm.?drain.*overflow',
        r'gutter.*overflow', r'drainage.?flood', r'culvert.?block',
        r'drain.?clog', r'clog.?drain', r'pluvial', r'rain.?flood',
        r'flood.?rain', r'monsoon.*flood', r'typhoon.*flood', r'cyclone.*flood',
    ]),
]

POLLUTION_SUBCATEGORY_RULES = [
    ('land_pollution', [
        r'land.?pollut', r'soil.?pollut', r'ground.?pollut', r'contaminat.*land',
        r'contaminat.*soil', r'oil.*spill.*land', r'chemical.*dump', r'landfill.*leach',
        r'toxic.*soil', r'heavy.?metal.*soil', r'industrial.*dump', r'waste.?land',
        r'illicit.*dump', r'open.?dump', r'municipal.?waste', r'biomedical.?waste',
        r'plastic.*land', r'plastic.*soil', r'land.?degrad', r'contamin.?site',
    ]),
    ('water_pollution', [
        r'water.?pollut', r'river.?pollut', r'lake.?pollut', r'pollut.*river',
        r'pollut.*water', r'pollut.*lake', r'pollut.*ocean', r'pollut.*sea',
        r'water.?contamin', r'sewage.*water', r'effluent', r'industrial.*water',
        r'wastewater', r'water.?quality.*poor', r'pollut.?canal', r'canal.?pollut',
        r'oil.*spill.*water', r'oil.*spill.*sea', r'oil.*spill.*ocean',
        r'chemical.*water', r'fertilizer.*runoff', r'pesticide.*runoff',
        r'runoff.*water', r'agricultural.*water', r'mining.*water',
        r'river.?dirty', r'river.?muddy.*pollut', r'estuary.?pollut',
        r'gulf.*pollut', r'bay.*pollut', r'harbour.?pollut', r'harbor.?pollut',
        r'foam.*river', r'foam.?pollut', r'industrial.?discharge', r'sludge.*water',
    ]),
    ('water_pollution_trash', [
        r'water.*trash', r'water.*garbage', r'river.*trash', r'river.*garbage',
        r'river.?plastic', r'plastic.?river', r'ocean.?plastic', r'sea.?plastic',
        r'plastic.*ocean', r'plastic.*sea', r'beach.*trash', r'beach.?garbage',
        r'marine.?plastic', r'marine.?debris', r'plastic.?debris.*water',
        r'floating.?trash', r'floating.?garbage', r'floating.?waste',
        r'canal.*trash', r'canal.?garbage', r'creek.*trash', r'creek.*garbage',
        r'creek.?plastic', r'ditch.?trash', r'ditch.?garbage', r'drain.?trash',
        r'bridge.*trash', r'stream.*trash',
    ]),
    ('air_pollution', [
        r'air.?pollut', r'pollut.*air', r'smog', r'smoke.*factory', r'factory.?smoke',
        r'emission', r'chimney.?smoke', r'chimney', r'power.?plant.*smoke',
        r'power.?plant.*pollut', r'vehicle.?emission', r'exhaust.*fume',
        r'fossil.?fuel', r'coal.?plant', r'air.?quality.*poor', r'pm.?2.?5',
        r'pm25', r'particulate', r'haze', r'industrial.?smoke',
        r'burning.*field', r'burning.?waste.*air', r'open.?burning',
        r'volcanic.?ash.*air', r'wildfire.*smoke', r'bushfire.*smoke',
        r'dust.?storm', r'cooking.?smoke.*pollut', r'incinerator',
    ]),
    ('electronics', [
        r'e.?waste', r'ewaste', r'electronic.?waste', r'electronic.?scrap',
        r'e.?scrap', r'circuit.?board', r'pcb', r'computer.*dump',
        r'monitor.*dump', r'tv.*dump', r'appliance.*dump', r'refriger.*dump',
        r'battery.*dump', r'lead.*battery', r'crt.*dump', r'cathode.?ray',
        r'electronic.?recycl', r'mobile.?phone.*dump', r'smartphone.*dump',
        r'cable.?dump', r'wire.?dump',
    ]),
]

NEG_SUBCATEGORY_RULES = [
    ('easy_negative', [
        r'portrait', r'selfie', r'person', r'face', r'family', r'wedding',
        r'birthday', r'party', r'cafe', r'restaurant', r'meal', r'food',
        r'drink', r'coffee', r'cat', r'dog', r'pet', r'animal', r'bird',
        r'flower', r'plant', r'garden', r'painting', r'sculpture', r'art',
        r'architecture', r'building', r'landscape.*no.*issue',
        r'clean.*park', r'clean.?street', r'clean.*beach', r'clean.?road',
        r'pristine', r'beautiful.*nature', r'scenic', r'sunset', r'sunrise',
        r'mountain.*clean', r'forest.*clean', r'beach.*clean', r'blue.?sky',
        r'sky.*cloud', r'abstract', r'pattern', r'texture', r'macro',
        r'closeup.*product', r'product.*shot', r'graphic', r'poster',
        r'logo', r'book.?cover', r'indoors', r'interior', r'furniture',
        r'toy', r'game', r'sport', r'soccer', r'basketball', r'tennis',
        r'hiking.*clean', r'camping.*clean', r'travel.*clean', r'vacation',
        r'concert', r'music.*festival.*clean', r'art.*gallery', r'museum',
        r'library', r'bookstore', r'clothing', r'fashion', r'cosmetic',
        r'makeup', r'nail', r'hairstyle', r'car.?clean', r'vehicle.*clean',
        r'transport.*clean', r'airplane.?clean', r'train.?clean', r'bus.?clean',
        r'ship.?clean', r'boat.*clean', r'education.*clean', r'student.*clean',
        r'classroom', r'office.*clean', r'meeting.*clean', r'conference.*clean',
        r'religious.*building', r'church', r'mosque', r'temple', r'synagogue',
        r'clean.*environment', r'clean.?city', r'clean.*area',
        r'healthy', r'fitness', r'exercise', r'yoga', r'meditation',
        r'news.?anchor', r'interview', r'studio.?shot', r'white.?background',
        r'product.?photography', r'still.?life',
    ]),
    ('hard_negative', [
        r'construction.*clean', r'building.*clean', r'road.?work.*clean',
        r'industrial.?clean', r'factory.*clean', r'warehouse.*clean',
        r'mining.*clean.*no.*pollut', r'agriculture.*clean', r'farm.*clean',
        r'rural.*clean', r'countryside.*clean', r'field.*clean',
        r'orchard.*clean', r'plantation.*clean', r'logging.*clean.*no.*damage',
        r'dam.*clean.*no.*flood', r'bridge.*clean', r'port.*clean',
        r'airport.*clean', r'station.*clean', r'terminal.*clean',
        r'hospital.*clean', r'clinic.*clean', r'restaurant.*clean',
        r'hotel.*clean', r'resort.*clean', r'mall.*clean', r'shop.*clean',
        r'store.*clean', r'supermarket.*clean', r'school.?clean',
        r'university.*clean', r'campus.*clean', r'stadium.*clean',
        r'arena.*clean', r'parking.?lot.*clean', r'playground.*clean',
        r'beach.*without.*trash', r'river.*clean.*no.*pollut',
        r'lake.*clean', r'waterfall.*clean', r'forest.*healthy',
        r'reef.*healthy', r'coral.*no.*bleach', r'wetland.*healthy',
        r'mountain.*hiking.*clean', r'desert.*clean', r'tundra.*clean',
        r'coast.*no.*erosion', r'shore.?stable',
        r'earthquake.?aftermath.*no.*env.*issue',
        r'fire.*but.?controlled.*no.*pollut', r'storm.?damage.*but.?cleanup',
    ]),
]


def classify_subcategory(prefix, title, description):
    combined = ('%s %s' % (title or '', description or '')).lower()

    if prefix == 'WST':
        rules = WASTE_SUBCATEGORY_RULES
        default = 'other'
    elif prefix == 'FLD':
        rules = FLOODING_SUBCATEGORY_RULES
        default = 'other'
    elif prefix == 'POL':
        rules = POLLUTION_SUBCATEGORY_RULES
        default = 'other'
    elif prefix == 'NEG':
        rules = NEG_SUBCATEGORY_RULES
        default = 'easy_negative'
    else:
        return None

    best_category = None
    best_score = 0
    for cat, patterns in rules:
        score = 0
        for pat in patterns:
            if re.search(pat, combined):
                score += 1
        if score > best_score:
            best_score = score
            best_category = cat

    if best_score > 0:
        return best_category
    return default


def is_acceptable_license(license_str):
    if not license_str:
        return False
    l = (license_str or '').lower()
    if 'all rights reserved' in l:
        return False
    if 'noncommercial' in l or 'non-commercial' in l or 'nc' in l:
        return False
    if 'noderivatives' in l or 'no-derivatives' in l or 'nd' in l:
        return False
    if 'creative commons' in l or 'public domain' in l or 'cc0' in l:
        return True
    return False


def resolution_str(width, height):
    try:
        w = int(width)
        h = int(height)
        if w <= 0 or h <= 0:
            return None
        return '{:,} \u256b {:,}'.format(w, h)
    except (ValueError, TypeError):
        return None


def quality_assessment(width, height, license_ok):
    try:
        w = int(width or 0)
        h = int(height or 0)
    except:
        w, h = 0, 0
    if w == 0 or h == 0:
        return None
    area = w * h
    if area < 300 * 300:
        return 'poor'
    if not license_ok:
        return 'usable'
    if area < 800 * 600:
        return 'fair'
    return 'good'


def main():
    with open(JSON_PATH, 'r', encoding='utf-8') as f:
        metadata = json.load(f)
    print('Loaded metadata for %d entries' % len(metadata))

    wb = load_workbook(EXCEL_PATH)
    ws = wb[SHEET_NAME]

    headers = [cell.value for cell in ws[1]]
    col_idx = {h: i + 1 for i, h in enumerate(headers)}
    print('Columns:', headers)

    rows_processed = 0
    rows_completed = 0
    rows_partial = 0
    inaccessible = 0

    for row_num in range(2, ws.max_row + 1):
        image_id = ws.cell(row=row_num, column=col_idx['image_id']).value
        if not image_id:
            continue

        prefix = image_id.split('_')[0]
        source_url = ws.cell(row=row_num, column=col_idx['source_url']).value
        if not source_url:
            continue

        meta = metadata.get(image_id, {})
        if meta.get('error'):
            inaccessible += 1

        # Track which cells were originally empty vs filled
        def get_val(col):
            v = ws.cell(row=row_num, column=col_idx[col]).value
            return v if (v is not None and v != '') else None

        def set_val_if_empty(col, new_val):
            if new_val is None or new_val == '':
                return False
            existing = get_val(col)
            if existing is not None:
                return False
            ws.cell(row=row_num, column=col_idx[col], value=new_val)
            return True

        any_change = False
        filled_important = 0
        total_important = 0

        # 1. filename
        if not get_val('filename'):
            total_important += 1
            if set_val_if_empty('filename', image_id + '.jpg'):
                any_change = True
                filled_important += 1

        # 2. primary_label
        if not get_val('primary_label'):
            total_important += 1
            if prefix in PRIMARY_LABEL_MAP:
                if set_val_if_empty('primary_label', PRIMARY_LABEL_MAP[prefix]):
                    any_change = True
                    filled_important += 1

        # 3. subcategory
        if not get_val('subcategory'):
            total_important += 1
            t = meta.get('title', '') or ''
            d = meta.get('description', '') or ''
            subcat = classify_subcategory(prefix, t, d)
            if set_val_if_empty('subcategory', subcat):
                any_change = True
                filled_important += 1

        # 4. difficulty
        if not get_val('difficulty'):
            total_important += 1
            # Default to clear per dataset convention; most images show clear subject
            if set_val_if_empty('difficulty', 'clear'):
                any_change = True
                filled_important += 1

        # 5. source_name
        if not get_val('source_name'):
            total_important += 1
            url_str = str(source_url).lower()
            if 'flickr.com' in url_str:
                sname = 'Flickr'
            elif 'wikimedia.org' in url_str or 'wikipedia.org' in url_str:
                sname = 'Wikimedia Commons'
            else:
                sname = None
            if set_val_if_empty('source_name', sname):
                any_change = True
                filled_important += 1

        # 6. original_author
        if not get_val('original_author'):
            total_important += 1
            author = meta.get('author_display') or meta.get('owner_username')
            if set_val_if_empty('original_author', author):
                any_change = True
                filled_important += 1

        # 7. license
        license_val = get_val('license')
        if not license_val:
            total_important += 1
            license_val = meta.get('license')
            if set_val_if_empty('license', license_val):
                any_change = True
                filled_important += 1

        # 8. license_url
        if not get_val('license_url'):
            total_important += 1
            if set_val_if_empty('license_url', meta.get('license_url')):
                any_change = True
                filled_important += 1

        # 9. attribution_require
        if not get_val('attribution_require'):
            total_important += 1
            if set_val_if_empty('attribution_require', meta.get('attribution_require')):
                any_change = True
                filled_important += 1

        # 10. download_date
        if not get_val('download_date'):
            total_important += 1
            if set_val_if_empty('download_date', DOWNLOAD_DATE):
                any_change = True
                filled_important += 1

        # 11. source_id
        if not get_val('source_id'):
            total_important += 1
            photo_id = meta.get('photo_id')
            if photo_id:
                sid_val = photo_id
            else:
                sid_val = '-'
            if set_val_if_empty('source_id', sid_val):
                any_change = True
                filled_important += 1

        # 12. resolution
        if not get_val('resolution'):
            total_important += 1
            w = meta.get('width')
            h = meta.get('height')
            rs = resolution_str(w, h)
            if set_val_if_empty('resolution', rs):
                any_change = True
                filled_important += 1

        # 13. quality_flag
        if not get_val('quality_flag'):
            total_important += 1
            license_ok = is_acceptable_license(license_val or meta.get('license', ''))
            w = meta.get('width')
            h = meta.get('height')
            qf = quality_assessment(w, h, license_ok)
            if set_val_if_empty('quality_flag', qf):
                any_change = True
                filled_important += 1

        # 14. location
        if not get_val('location'):
            total_important += 1
            loc_hint = meta.get('location_hint')
            if not loc_hint:
                d = meta.get('description') or ''
                t = meta.get('title') or ''
                combined = t + ' ' + d
                # Look for Jakarta/Bangkok/etc. patterns
                city_pats = [
                    (r'jakarta', 'Jakarta, Indonesia'),
                    (r'banjir', 'Jakarta, Indonesia'),
                    (r'bangkok', 'Bangkok, Thailand'),
                    (r'thailand.*flood', 'Thailand'),
                    (r'flood.*thailand', 'Thailand'),
                    (r'chumpon', 'Chumphon, Thailand'),
                    (r'surat.?thani', 'Surat Thani, Thailand'),
                    (r'penang', 'Penang, Malaysia'),
                    (r'sungai.?pinang', 'Penang, Malaysia'),
                    (r'vietnam.*flood', 'Vietnam'),
                    (r'kampung', 'Indonesia'),
                    (r'kemang', 'Jakarta, Indonesia'),
                    (r'kapuk.?muara', 'Jakarta, Indonesia'),
                    (r'jl.?haji', 'Jakarta, Indonesia'),
                    (r'pos.?pengumben', 'Jakarta, Indonesia'),
                    (r'h.?rasuna.?said', 'Jakarta, Indonesia'),
                    (r'chatuchak', 'Bangkok, Thailand'),
                    (r'patumtani', 'Thailand'),
                    (r'khao.?san.?road', 'Bangkok, Thailand'),
                    (r'goa.?india', 'Goa, India'),
                    (r'cortalim', 'Goa, India'),
                    (r'panjim', 'Panaji, Goa, India'),
                    (r'campal', 'Campal, Goa, India'),
                    (r'd.?b.?road.*goa', 'Goa, India'),
                    (r'nh17.*goa', 'Goa, India'),
                    (r'bhandeam', 'Goa, India'),
                    (r'goel.?ganga', 'Pune, India'),
                    (r'pune', 'Pune, India'),
                    (r'karnataka', 'Karnataka, India'),
                    (r'maligaya.?bridge', 'Philippines'),
                    (r'marilao', 'Marilao, Bulacan, Philippines'),
                    (r'santa.?maria.*macabebe', 'Pampanga, Philippines'),
                    (r'dagupan', 'Dagupan, Philippines'),
                    (r'kathmandu', 'Kathmandu, Nepal'),
                    (r'indonesia', 'Indonesia'),
                    (r'philippines', 'Philippines'),
                    (r'malaysia', 'Malaysia'),
                    (r'south.?asia', 'South Asia'),
                    (r'cyclone.?roanu', 'Bangladesh/Sri Lanka/India'),
                ]
                for pat, name in city_pats:
                    if re.search(pat, combined, re.I):
                        loc_hint = name
                        break
            if set_val_if_empty('location', loc_hint):
                any_change = True
                filled_important += 1

        # 15. event_group
        if not get_val('event_group'):
            t = meta.get('title') or ''
            d = meta.get('description') or ''
            combined = t + ' ' + d
            eg = None
            # Event groups for obvious clustered events
            if re.search(r'thailand.*flood.*2011|2011.*thailand.*flood|thai.*flood.*2011|bangkok.*flood.*2011', combined, re.I):
                eg = 'Thailand 2011 Floods'
            elif re.search(r'jakarta.*flood|banjir.*jakarta', combined, re.I):
                eg = 'Jakarta Floods'
            elif re.search(r'bin.?strike', combined, re.I):
                eg = 'London Bin Strike 2015'
            elif re.search(r'typhoon.?goni', combined, re.I):
                eg = 'Typhoon Goni 2020'
            elif re.search(r'cyclone.?roanu', combined, re.I):
                eg = 'Cyclone Roanu 2016'
            elif re.search(r'cortalim|panjim|campal|d.?b.?road|nh17|goa.*roadside|dirtypanjim|joegoauk', combined, re.I):
                eg = 'Goa Roadside Waste Survey'
            elif re.search(r'h.?rasuna.?said|pos.?pengumben|kapuk.?muara|kemang|jl.?haji', combined, re.I):
                eg = 'Jakarta Urban Flooding'
            if eg:
                total_important += 1
                if set_val_if_empty('event_group', eg):
                    any_change = True
                    filled_important += 1
            else:
                set_val_if_empty('event_group', '-')

        # 16. notes
        if not get_val('notes'):
            t = meta.get('title') or ''
            d = meta.get('description') or ''
            note = None
            if prefix == 'NEG':
                note = 'Non-environmental subject: %s.' % (t or 'Generic subject.')
            elif t and len(t) > 5 and t.lower() not in ('garbage', 'trash', 'rubbish', 'litter', 'floods', 'flooded', 'flood'):
                short = t[:80].strip()
                if not re.match(r'^(IMG|DSC|DSCF|P|102 copy|ESL|0\d{2}-|flood-\d|rain-\d)', short, re.I):
                    note = short + ('.' if not short.endswith('.') else '')
            if d and not note:
                short_d = d[:80].strip()
                if short_d and len(short_d) > 10:
                    note = short_d
            if meta.get('error'):
                note = 'Source error: %s' % meta['error'][:60]
            if note:
                set_val_if_empty('notes', note)

        # 17. license_verified
        if not get_val('license_verified'):
            total_important += 1
            lic = get_val('license') or license_val
            lic_url = get_val('license_url') or meta.get('license_url')
            if lic and lic_url and 'error' not in meta:
                if set_val_if_empty('license_verified', 'Yes'):
                    any_change = True
                    filled_important += 1

        # 18. approved_for_training
        if not get_val('approved_for_training'):
            total_important += 1
            lic = get_val('license') or license_val
            qf = get_val('quality_flag') or meta.get('quality_flag')
            lic_ok = is_acceptable_license(lic)
            q_ok = qf in ('good', 'fair', 'usable')
            sub_ok = get_val('subcategory') is not None
            if lic_ok and q_ok and sub_ok:
                if set_val_if_empty('approved_for_training', 'Yes'):
                    any_change = True
                    filled_important += 1

        if any_change:
            rows_processed += 1
            if total_important > 0 and filled_important == total_important:
                rows_completed += 1
            else:
                rows_partial += 1

    wb.save(EXCEL_PATH)
    print('\n=== UPDATE SUMMARY ===')
    print('Rows with changes:', rows_processed)
    print('  - Fully completed:', rows_completed)
    print('  - Partially completed:', rows_partial)
    print('Rows with scrape errors:', inaccessible)
    print('Saved to:', EXCEL_PATH)


if __name__ == '__main__':
    main()
