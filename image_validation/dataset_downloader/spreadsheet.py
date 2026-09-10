"""
spreadsheet.py -- Read and parse the EcoPin dataset spreadsheet.

Determines which rows are PENDING, ALREADY_DONE, or SKIP based on
the actual Excel cell fill colors (not just data presence).
"""

from __future__ import annotations

import os
from dataclasses import dataclass, field
from enum import Enum
from typing import Optional

import openpyxl
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

import config


class RowStatus(str, Enum):
    PENDING            = "pending"
    ALREADY_EXISTS     = "already_exists"
    SKIP_PARTIAL       = "skip_partial"
    NO_URL             = "no_url"
    UNKNOWN_SOURCE     = "unknown_source"
    UNKNOWN_LABEL      = "unknown_label"


@dataclass
class ImageRow:
    """Represents one image entry from the Dataset sheet."""
    excel_row:          int
    image_id:           str
    filename:           str
    primary_label:      str
    subcategory:        str
    difficulty:         str
    source_name:        str
    source_url:         str
    original_author:    str
    license:            str
    license_url:        str
    attribution_req:    str
    download_date:      str
    notes:              str
    # Derived
    status:             RowStatus = RowStatus.PENDING
    destination_folder: str = ""
    destination_path:   str = ""
    yellow_cols:        list = field(default_factory=list)
    green_col1:         bool = False
    on_disk:            bool = False


def _get_cell_rgb(cell) -> Optional[str]:
    """Return the ARGB hex string of a cell's foreground fill color, or None."""
    fill = cell.fill
    if fill is None:
        return None
    if fill.fill_type in (None, "none"):
        return None
    fg = fill.fgColor
    if fg is None:
        return None
    if fg.type == "rgb":
        rgb = fg.rgb
        if rgb in ("00000000", "FFFFFFFF"):
            return None
        return rgb
    return None


def _collect_fill_colors(ws: Worksheet, row_num: int) -> tuple[list[int], list[int]]:
    """
    Return (yellow_cols, green_cols) -- lists of 1-based column indices
    that have the corresponding highlight color.
    """
    yellow_cols: list[int] = []
    green_cols:  list[int] = []
    for col_idx in range(1, config.TOTAL_COLS + 1):
        rgb = _get_cell_rgb(ws.cell(row=row_num, column=col_idx))
        if rgb == config.YELLOW_COLOR:
            yellow_cols.append(col_idx)
        elif rgb == config.PENDING_GREEN_COLOR:
            green_cols.append(col_idx)
    return yellow_cols, green_cols


def _destination_path(label: str, filename: str) -> tuple[str, str]:
    """
    Return (folder_path, full_file_path) for the given label and filename.
    Returns ("", "") if the label is not in LABEL_FOLDER_MAP.
    """
    folder_name = config.LABEL_FOLDER_MAP.get(label, "")
    if not folder_name:
        return "", ""
    folder_path = os.path.join(config.DATASET_RAW_DIR, folder_name)
    file_path   = os.path.join(folder_path, filename)
    return folder_path, file_path


def _classify_row(
    yellow_cols: list[int],
    green_cols:  list[int],
    on_disk:     bool,
) -> RowStatus:
    """
    Apply the highlighting rules to determine row status.

    Rules (in priority order):
    1. Green on col 1  ->  PENDING
    2. Only col 20 yellow  ->  ALREADY_EXISTS (approved_for_training)
    3. Many yellow cols (≥ MIN_YELLOW_COLS) + NOT on disk  ->  PENDING
    4. Many yellow cols + already on disk  ->  ALREADY_EXISTS
    5. Everything else  ->  SKIP_PARTIAL
    """
    has_green_col1 = 1 in green_cols

    if has_green_col1:
        return RowStatus.PENDING

    if yellow_cols == [config.COL_APPROVED_TRAINING]:
        return RowStatus.ALREADY_EXISTS

    if len(yellow_cols) >= config.MIN_YELLOW_COLS_FOR_PENDING:
        if on_disk:
            return RowStatus.ALREADY_EXISTS
        return RowStatus.PENDING

    return RowStatus.SKIP_PARTIAL


def load_rows(diagnostic: bool = False) -> list[ImageRow]:
    """
    Open the spreadsheet and return a list of ImageRow objects,
    each with a computed status.

    If diagnostic=True, print a per-row status summary.
    """
    print(f"Loading spreadsheet: {config.EXCEL_FILE}")
    wb: Workbook = openpyxl.load_workbook(config.EXCEL_FILE)

    if config.DATASET_SHEET not in wb.sheetnames:
        raise ValueError(
            f"Sheet '{config.DATASET_SHEET}' not found. "
            f"Available: {wb.sheetnames}"
        )

    ws: Worksheet = wb[config.DATASET_SHEET]
    rows: list[ImageRow] = []

    for row_num in range(2, ws.max_row + 1):
        image_id = ws.cell(row=row_num, column=config.COL_IMAGE_ID).value
        if not image_id:
            continue  # blank row

        def _str(col: int) -> str:
            v = ws.cell(row=row_num, column=col).value
            return str(v).strip() if v is not None else ""

        filename      = _str(config.COL_FILENAME)
        primary_label = _str(config.COL_PRIMARY_LABEL)
        subcategory   = _str(config.COL_SUBCATEGORY)
        difficulty    = _str(config.COL_DIFFICULTY)
        source_name   = _str(config.COL_SOURCE_NAME)
        source_url    = _str(config.COL_SOURCE_URL)
        author        = _str(config.COL_ORIGINAL_AUTHOR)
        license_      = _str(config.COL_LICENSE)
        license_url   = _str(config.COL_LICENSE_URL)
        attr_req      = _str(config.COL_ATTRIBUTION_REQUIRE)
        dl_date       = _str(config.COL_DOWNLOAD_DATE)
        notes         = _str(config.COL_NOTES)

        # Resolve destination
        dest_folder, dest_path = _destination_path(primary_label, filename)

        # Check disk
        on_disk = os.path.isfile(dest_path) if dest_path else False

        # Fill analysis
        yellow_cols, green_cols = _collect_fill_colors(ws, row_num)

        # Classify
        if not source_url:
            status = RowStatus.NO_URL
        elif source_name not in (config.SOURCE_WIKIMEDIA, config.SOURCE_FLICKR):
            status = RowStatus.UNKNOWN_SOURCE
        elif not dest_folder:
            status = RowStatus.UNKNOWN_LABEL
        else:
            status = _classify_row(yellow_cols, green_cols, on_disk)

        image_row = ImageRow(
            excel_row          = row_num,
            image_id           = str(image_id),
            filename           = filename,
            primary_label      = primary_label,
            subcategory        = subcategory,
            difficulty         = difficulty,
            source_name        = source_name,
            source_url         = source_url,
            original_author    = author,
            license            = license_,
            license_url        = license_url,
            attribution_req    = attr_req,
            download_date      = dl_date,
            notes              = notes,
            status             = status,
            destination_folder = dest_folder,
            destination_path   = dest_path,
            yellow_cols        = yellow_cols,
            green_col1         = (1 in green_cols),
            on_disk            = on_disk,
        )
        rows.append(image_row)

        if diagnostic:
            _print_diagnostic(image_row)

    wb.close()
    return rows


def _print_diagnostic(r: ImageRow) -> None:
    green_marker = " [GREEN-A]" if r.green_col1 else ""
    disk_marker  = " [ON DISK]" if r.on_disk else ""
    print(
        f"  Row {r.excel_row:4d}  {r.image_id:<12}  "
        f"yellow={len(r.yellow_cols):2d}cols{green_marker}{disk_marker}  "
        f"-> {r.status.value.upper()}"
    )


def summarise(rows: list[ImageRow]) -> None:
    """Print a count summary by status."""
    from collections import Counter
    counts = Counter(r.status for r in rows)
    total  = len(rows)
    print()
    print("-" * 50)
    print(f"  Total rows parsed:  {total}")
    for status in RowStatus:
        print(f"  {status.value:<30} {counts[status]:4d}")
    print("-" * 50)
    print()
